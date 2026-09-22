# -*- coding: utf-8 -*-
"""990 store - sample DOCX generators (resume / cover letter / IOU / quote / minutes)"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r"C:\Users\iss59\Desktop\antigravity\990store\files"
FONT = "맑은 고딕"


def set_cell_bg(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def set_font(run, size=10, bold=False, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def para(doc, text="", size=10, bold=False, align=None, space_after=4, space_before=0, color=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    if text:
        set_font(p.add_run(text), size, bold, color)
    return p


def title(doc, text, size=18):
    p = para(doc, text, size=size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2, space_before=0)
    # thin line under title
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)
    pPr = p2._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:color"), "4472C4")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def fill(cell, text, bold=False, size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT, bg=None, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    set_font(p.add_run(text), size, bold, color)
    if bg:
        set_cell_bg(cell, bg)


def header_row(table, labels, bg="4472C4", fg=(255, 255, 255)):
    for i, label in enumerate(labels):
        fill(table.rows[0].cells[i], label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, bg=bg, color=fg)


def style_table(table):
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER


def set_col_widths(table, widths_cm):
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            if i < len(row.cells):
                row.cells[i].width = Cm(w)


def label_table(doc, rows):
    """rows: list of (label, value) -> 2-col table, label cells shaded"""
    t = doc.add_table(rows=len(rows), cols=2)
    style_table(t)
    set_col_widths(t, [3.2, 13.0])
    for i, (lab, val) in enumerate(rows):
        fill(t.rows[i].cells[0], lab, bold=True, bg="D9E2F3")
        fill(t.rows[i].cells[1], val)
    return t


def make_resume():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)

    title(doc, "이 력 서")

    # top: photo + basic info
    t = doc.add_table(rows=4, cols=4)
    style_table(t)
    set_col_widths(t, [2.6, 5.6, 2.6, 5.4])
    fill(t.rows[0].cells[0], "이름", True, bg="D9E2F3", align=WD_ALIGN_PARAGRAPH.CENTER)
    fill(t.rows[0].cells[1], "")
    fill(t.rows[0].cells[2], "생년월일", True, bg="D9E2F3", align=WD_ALIGN_PARAGRAPH.CENTER)
    fill(t.rows[0].cells[3], "")
    fill(t.rows[1].cells[0], "연락처", True, bg="D9E2F3", align=WD_ALIGN_PARAGRAPH.CENTER)
    fill(t.rows[1].cells[1], "")
    fill(t.rows[1].cells[2], "이메일", True, bg="D9E2F3", align=WD_ALIGN_PARAGRAPH.CENTER)
    fill(t.rows[1].cells[3], "")
    fill(t.rows[2].cells[0], "주소", True, bg="D9E2F3", align=WD_ALIGN_PARAGRAPH.CENTER)
    merged = t.rows[2].cells[1].merge(t.rows[2].cells[3])
    fill(merged, "")
    photo = t.rows[3].cells[0].merge(t.rows[3].cells[1])
    fill(photo, "사 진 (3.5cm × 4.5cm)", align=WD_ALIGN_PARAGRAPH.CENTER, color=(150, 150, 150))
    qual = t.rows[3].cells[2].merge(t.rows[3].cells[3])
    fill(qual, "지원 직무 : ", bold=True)

    para(doc, "", space_after=4)

    # education
    para(doc, "학 력", size=11, bold=True, space_before=4, space_after=4, color=(0x44, 0x72, 0xC4))
    t = doc.add_table(rows=4, cols=4)
    style_table(t)
    set_col_widths(t, [3.4, 5.0, 5.0, 2.8])
    header_row(t, ["기간", "학교명", "전공 / 내용", "학점"])
    for r in range(1, 4):
        for c in range(4):
            fill(t.rows[r].cells[c], "")

    para(doc, "", space_after=4)

    # career
    para(doc, "경 력", size=11, bold=True, space_before=4, space_after=4, color=(0x44, 0x72, 0xC4))
    t = doc.add_table(rows=4, cols=4)
    style_table(t)
    set_col_widths(t, [3.4, 4.0, 6.4, 2.4])
    header_row(t, ["기간", "회사명", "담당 업무", "퇴사 사유"])
    for r in range(1, 4):
        for c in range(4):
            fill(t.rows[r].cells[c], "")

    para(doc, "", space_after=4)

    # certificates + skills side info
    para(doc, "자격증 및 어학", size=11, bold=True, space_before=4, space_after=4, color=(0x44, 0x72, 0xC4))
    t = doc.add_table(rows=3, cols=4)
    style_table(t)
    set_col_widths(t, [3.4, 5.0, 4.8, 3.0])
    header_row(t, ["취득일", "자격증명", "발급기관", "점수/급수"])
    for r in range(1, 3):
        for c in range(4):
            fill(t.rows[r].cells[c], "")

    para(doc, "", space_after=4)

    # skills
    para(doc, "보유 기술 / 특기", size=11, bold=True, space_before=4, space_after=4, color=(0x44, 0x72, 0xC4))
    t = doc.add_table(rows=1, cols=1)
    style_table(t)
    fill(t.rows[0].cells[0], "· 예) 한글/엑셀/파워포인트 (활용 수준 기재)\n· 예) 프로그램 스킬 : Photoshop, Figma, Python 등\n· 예) 어학 : TOEIC 850 (2025.03), OPIc IH 등")

    para(doc, "", space_after=4)

    # self intro
    para(doc, "자기소개", size=11, bold=True, space_before=4, space_after=4, color=(0x44, 0x72, 0xC4))
    t = doc.add_table(rows=1, cols=1)
    style_table(t)
    fill(t.rows[0].cells[0],
         " (지원 직무와 연결되는 경험 1~2개를 구체적 수치와 함께 5~7줄로 작성하세요. "
         "'무엇을 배우고, 어떤 행동을 했고, 어떤 결과를 냈는지' 구조를 유지하면 좋습니다.)",
         color=(120, 120, 120))

    para(doc, "", space_after=6)
    p = para(doc, "작성일 : 2026 년    월    일        지원자 :                    (인) ",
             size=10, align=WD_ALIGN_PARAGRAPH.RIGHT)

    path = os.path.join(BASE, "resume", "이력서양식_표준형.docx")
    doc.save(path)
    return path


def make_cover_letter():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

    title(doc, "자기소개서")

    label_table(doc, [
        ("지원 회사", ""),
        ("지원 직무", ""),
        ("이 립 일", "2026 년    월    일"),
        ("작성 예시", "아래는 합격 사례를 구조화한 예시입니다. [ ] 부분을 본인 경험으로 바꾸세요."),
    ])

    para(doc, "", space_after=6)

    sections = [
        ("1. 성장 과정",
         "[가족/환경]에서 자라며 [무엇을 자연스럽게 배웠는지] 써 주세요.\n"
         "예시) 형제 사이에서 중재자 역할을 자주 맡으며, 서로 다른 의견을 조율하는 것에 익숙해졌습니다. "
         "대학 동아리에서도 팀원 간 갈등을 먼저 해결해 프로젝트 일정을 2주 앞당긴 경험이 있습니다.\n\n"
         "작성 팁 : 성장 과정은 '지원 직무에 필요한 역량'으로 이어지게 끝맺으세요."),
        ("2. 성격의 장·단점",
         "장점 : [상황]에서 [행동]하여 [결과] (구체적 사례 1개)\n"
         "단점 : [단점]을 인정 + 이를 보완하기 위한 [노력/시스템]\n\n"
         "예시 - 장점) 마감 3일 전 자료가 미비한 상태에서, 표준 템플릿을 만들어 팀 5명의 작성 속도를 2배 높였습니다.\n"
         "예시 - 단점) 세부 확인에 시간이 오래 걸리는 편이라, 마감 2일 전 자동 체크리스트를 만들어 실수를 줄였습니다.\n\n"
         "작성 팁 : 단점은 '변명'이 아니라 '개선 노력'이 중심이어야 합니다."),
        ("3. 지원 동기",
         "[회사의 제품/문화/공시 등 구체적 이유] + [직무와의 연결] + [입사 후 기여 포인트]\n\n"
         "예시) ○○의 ○○ 서비스를 6개월간 사용하며 [구체적 경험]을 했고, 이 과정에서 [개선점]을 발견했습니다. "
         "입사 후 데이터 기반으로 해당 프로세스를 개선하고 싶습니다.\n\n"
         "작성 팁 : 회사 공식 자료(IR, 블로그, 채용공고)에서 사실을 1개 이상 꼭 넣으세요."),
        ("4. 입사 후 포부",
         "[1단계 : 적응/학습 3개월] → [2단계 : 업무 기여 1년] → [3단계 : 장기 목표]\n\n"
         "예시) 첫 3개월은 사내 데이터 구조와 프로세스 학습에 집중하겠습니다. 1년 내에는 담당 채널의 "
         "전환율을 [목표치]까지 끌어올리는 것을 목표로 하겠습니다.\n\n"
         "작성 팁 : '열심히 하겠습니다' 같은 추상 표현 대신 숫자 목표를 넣으면 신뢰도가 올라갑니다."),
    ]

    for head, body in sections:
        para(doc, head, size=12, bold=True, space_before=8, space_after=4, color=(0x44, 0x72, 0xC4))
        t = doc.add_table(rows=1, cols=1)
        style_table(t)
        fill(t.rows[0].cells[0], body, size=10)
        para(doc, "", space_after=4)

    para(doc, "※ 자소서 체크리스트 : □ 수치 포함 □ 직무 연관성 □ 사실 확인 □ 맞춤법 검사 □ 분량 준수(각 항목 600~800자)",
         size=9, color=(150, 60, 60), space_before=8)

    path = os.path.join(BASE, "resume", "자기소개서_작성예시_4항목.docx")
    doc.save(path)
    return path


def make_iou():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)

    title(doc, "차 용 증", size=20)
    para(doc, "(금전소비대차계약서)", size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=(120, 120, 120), space_after=10)

    para(doc, "채무자(차주)와 금전대차인(대주)은 다음과 같이 금전소비대차계약을 체결한다.",
         size=10.5, space_after=8)

    label_table(doc, [
        ("차 주 (채무자)", "성명 :                    (서명 :            )    주민등록번호 :                    -"),
        ("연락처", "연락처 :                                 주소 :"),
        ("금전대차인 (대주)", "성명 :                    (서명 :            )    주민등록번호 :                    -"),
        ("연락처", "연락처 :                                 주소 :"),
    ])

    para(doc, "", space_after=6)

    para(doc, "제1조 (금전의 지급) 대주는 채무자에게 다음과 같은 조건으로 금전을 대차한다.",
         size=10.5, bold=True, space_before=4, space_after=4)

    label_table(doc, [
        ("대차 금액", "금                          원整 (₩                    원)"),
        ("이 자 율", "연         %   (법정 최고이자율 이내)   ※ 무이자인 경우 '무이자' 기재"),
        ("변제 기일", "2026 년    월    일"),
        ("지급 방법", "□ 계좌이체   □ 현금 직접 지급   □ 기타 :"),
    ])

    para(doc, "", space_after=6)

    clauses = [
        ("제2조 (지연이자)", "채무자가 제1조의 변제기일까지 변제하지 아니하는 때에는 연        %의 비율로 계산된 지연손해금을 지급하여야 한다. (법정 최고이자율 이내로 한정)"),
        ("제3조 (기한의 이익 상실)", "① 채무자가 이자 또는 변제금을 정해진 기일에 2회 이상 지급하지 아니한 때\n② 채무자가 강제집행, 회생절차, 파산선고 등을 받은 때\n③ 담보 가치가 현저히 감소한 때\n→ 대주는 기한의 이익을 상실시켜 잔여 원금의 즉시 변제를 청구할 수 있다."),
        ("제4조 (계약의 해지)", "당사자는 상대방에 대한 서면(문자 포함) 통지로 본 계약을 해지할 수 있다. 단, 이미 발생한 채무는 소멸하지 않는다."),
    ]
    for head, body in clauses:
        para(doc, head, size=10.5, bold=True, space_before=6, space_after=2)
        para(doc, body, size=10, space_after=4)

    para(doc, "", space_after=8)
    para(doc, "2026 년          월          일", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    t = doc.add_table(rows=1, cols=2)
    style_table(t)
    set_col_widths(t, [8.2, 8.2])
    fill(t.rows[0].cells[0], "금전대차인 (대주) :              (인)", align=WD_ALIGN_PARAGRAPH.CENTER)
    fill(t.rows[0].cells[1], "차 주 (채무자) :              (인)", align=WD_ALIGN_PARAGRAPH.CENTER)

    para(doc, "", space_after=6)
    para(doc, "※ 본 서식은 참고용 샘플이며, 구체적인 거래는 법률 전문가 상담을 권장합니다. 금전소비대차계약은 대주·채무자 간 합의로 성립하며, 이자는 법정 최고이자율(연 20%)을 초과할 수 없습니다.",
         size=8.5, color=(150, 60, 60), space_before=6)

    path = os.path.join(BASE, "life", "차용증_금전소비대차_서식.docx")
    doc.save(path)
    return path


def make_quote():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)

    title(doc, "견 적 서")

    para(doc, "견적번호 : QT-2026-______        견적일 : 2026 년    월    일        유효기간 : 견적일로부터 14일",
         size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8, color=(80, 80, 80))

    # supplier / buyer
    t = doc.add_table(rows=6, cols=4)
    style_table(t)
    set_col_widths(t, [2.4, 5.8, 2.4, 5.6])

    fill(t.rows[0].cells[0], "공 급 자", True, bg="4472C4", color=(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)
    m = t.rows[0].cells[1].merge(t.rows[0].cells[3])
    fill(m, "")
    for i, (lab, lab2) in enumerate([("사업자명", "사업자등록번호"), ("대 표 자", "연락처"), ("이 메 일", "주소")], start=1):
        fill(t.rows[i].cells[0], lab, True, bg="D9E2F3", align=WD_ALIGN_PARAGRAPH.CENTER)
        fill(t.rows[i].cells[1], "")
        fill(t.rows[i].cells[2], lab2, True, bg="D9E2F3", align=WD_ALIGN_PARAGRAPH.CENTER)
        fill(t.rows[i].cells[3], "")

    fill(t.rows[5].cells[0], "공급받는자", True, bg="4472C4", color=(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)
    m = t.rows[5].cells[1].merge(t.rows[5].cells[3])
    fill(m, "담당자 :                       회사명 :                       연락처 :")

    para(doc, "", space_after=6)

    # items
    para(doc, "견적 내역", size=11, bold=True, space_before=4, space_after=4, color=(0x44, 0x72, 0xC4))
    t = doc.add_table(rows=8, cols=6)
    style_table(t)
    set_col_widths(t, [1.2, 5.6, 3.0, 2.0, 2.6, 2.6])
    header_row(t, ["No", "품목 / 용역명", "규격 / 내용", "수량", "단가", "금액"])
    for r in range(1, 8):
        fill(t.rows[r].cells[0], str(r), align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in range(1, 6):
            fill(t.rows[r].cells[c], "", align=WD_ALIGN_PARAGRAPH.RIGHT if c >= 4 else WD_ALIGN_PARAGRAPH.LEFT)

    para(doc, "", space_after=4)

    # totals
    t = doc.add_table(rows=4, cols=2)
    style_table(t)
    set_col_widths(t, [11.0, 5.4])
    totals = [("공급가액 (합계)", ""), ("부가세 (10%)", ""), ("총 계 (VAT 포함)", "")]
    for i, (lab, val) in enumerate(totals):
        bg = "D9E2F3" if i < 2 else "4472C4"
        fg = None if i < 2 else (255, 255, 255)
        fill(t.rows[i].cells[0], lab, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, bg=bg, color=fg)
        fill(t.rows[i].cells[1], val, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, bg=bg, color=fg)
    fill(t.rows[3].cells[0], "결제 조건", True, bg="D9E2F3")
    fill(t.rows[3].cells[1], "□ 선금 50%  □ 후금  □ 기타 :")

    para(doc, "", space_after=6)
    para(doc, "특이사항", size=11, bold=True, space_before=4, space_after=4, color=(0x44, 0x72, 0xC4))
    t = doc.add_table(rows=1, cols=1)
    style_table(t)
    fill(t.rows[0].cells[0],
         "· 납기 : 발주 확정 후    일 이내\n· 배송/설치 : 포함 / 미포함 (삭제)\n· A/S 및 환불 정책 :\n· 기타 조건 :")

    para(doc, "", space_after=8)
    p = para(doc, "위와 같이 견적합니다.", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    para(doc, "2026 년    월    일", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    t = doc.add_table(rows=1, cols=1)
    style_table(t)
    fill(t.rows[0].cells[0], "공급자 :  (주)                    대표자 :                    (인)", align=WD_ALIGN_PARAGRAPH.CENTER)

    path = os.path.join(BASE, "life", "견적서_표준_양식.docx")
    doc.save(path)
    return path


def make_minutes():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)

    title(doc, "회 의 록")

    label_table(doc, [
        ("회의명", ""),
        ("일 시", "2026 년    월    일 (   )    :     ~    :"),
        ("장 소", ""),
        ("참석자", ""),
        ("작성자", ""),
    ])

    para(doc, "", space_after=6)

    # agenda 1
    para(doc, "1. 안건", size=12, bold=True, space_before=4, space_after=4, color=(0x44, 0x72, 0xC4))
    t = doc.add_table(rows=4, cols=4)
    style_table(t)
    set_col_widths(t, [1.2, 5.0, 7.4, 2.8])
    header_row(t, ["No", "안건명", "논의 내용", "소요시간"])
    for r in range(1, 4):
        fill(t.rows[r].cells[0], str(r), align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in range(1, 4):
            fill(t.rows[r].cells[c], "")

    para(doc, "", space_after=6)

    # decisions
    para(doc, "2. 결정사항", size=12, bold=True, space_before=4, space_after=4, color=(0x44, 0x72, 0xC4))
    t = doc.add_table(rows=3, cols=3)
    style_table(t)
    set_col_widths(t, [1.2, 12.0, 3.2])
    header_row(t, ["No", "결정 내용", "의사결정자"])
    for r in range(1, 3):
        fill(t.rows[r].cells[0], str(r), align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in (1, 2):
            fill(t.rows[r].cells[c], "")

    para(doc, "", space_after=6)

    # todo
    para(doc, "3. ToDo (액션 아이템)", size=12, bold=True, space_before=4, space_after=4, color=(0x44, 0x72, 0xC4))
    t = doc.add_table(rows=5, cols=5)
    style_table(t)
    set_col_widths(t, [1.0, 6.6, 3.0, 3.0, 2.8])
    header_row(t, ["No", "할 일", "담당", "기한", "상태"])
    todos = ["", "", "", ""]
    for r in range(1, 5):
        fill(t.rows[r].cells[0], str(r), align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in range(1, 4):
            fill(t.rows[r].cells[c], "")
        fill(t.rows[r].cells[4], "미완", align=WD_ALIGN_PARAGRAPH.CENTER, bg="FFF2CC")

    para(doc, "", space_after=6)

    # next meeting + free notes
    para(doc, "4. 다음 회의 / 자유 기록", size=12, bold=True, space_before=4, space_after=4, color=(0x44, 0x72, 0xC4))
    t = doc.add_table(rows=2, cols=1)
    style_table(t)
    fill(t.rows[0].cells[0], "다음 회의 : 2026 년    월    일 (   )    :     / 장소 :")
    fill(t.rows[1].cells[0], "자유 기록 :\n")

    path = os.path.join(BASE, "office", "회의록_표준_양식.docx")
    doc.save(path)
    return path


if __name__ == "__main__":
    paths = [make_resume(), make_cover_letter(), make_iou(), make_quote(), make_minutes()]
    for p in paths:
        print("OK:", p)
