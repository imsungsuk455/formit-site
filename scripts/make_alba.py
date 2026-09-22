# -*- coding: utf-8 -*-
"""알바이력서 미끼 무료 상품 생성"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from make_samples import (
    BASE, title, para, fill, header_row, style_table, set_col_widths,
    label_table, set_font,
)


def make_alba_resume():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.5)
    sec.bottom_margin = Cm(1.5)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)

    title(doc, "알바 이력서")

    para(doc, "※ 편의점·카페·식당·과외 등 단기 알바 지원용으로 바로 쓰는 표준 서식입니다.",
         size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=(120, 120, 120), space_after=8)

    # basic info
    t = doc.add_table(rows=3, cols=4)
    style_table(t)
    set_col_widths(t, [2.4, 5.6, 2.4, 5.6])
    fill(t.rows[0].cells[0], "이름", True, bg="D9E2F3", align=WD_ALIGN_PARAGRAPH.CENTER)
    fill(t.rows[0].cells[1], "")
    fill(t.rows[0].cells[2], "생년월일", True, bg="D9E2F3", align=WD_ALIGN_PARAGRAPH.CENTER)
    fill(t.rows[0].cells[3], "")
    fill(t.rows[1].cells[0], "연락처", True, bg="D9E2F3", align=WD_ALIGN_PARAGRAPH.CENTER)
    fill(t.rows[1].cells[1], "")
    fill(t.rows[1].cells[2], "이메일", True, bg="D9E2F3", align=WD_ALIGN_PARAGRAPH.CENTER)
    fill(t.rows[1].cells[3], "")
    fill(t.rows[2].cells[0], "거주지", True, bg="D9E2F3", align=WD_ALIGN_PARAGRAPH.CENTER)
    m = t.rows[2].cells[1].merge(t.rows[2].cells[3])
    fill(m, "")

    para(doc, "", space_after=4)

    # availability - the key alba section
    para(doc, "근무 가능 시간", size=11, bold=True, space_before=4, space_after=4, color=(0x44, 0x72, 0xC4))
    t = doc.add_table(rows=4, cols=5)
    style_table(t)
    set_col_widths(t, [3.6, 2.7, 2.7, 2.7, 2.7])
    header_row(t, ["구분", "오전 (09~12)", "오후 (12~18)", "저녁 (18~22)", "심야 (22~)"])
    days = ["평일 (월~금)", "토요일", "일요일"]
    for i, d in enumerate(days, start=1):
        fill(t.rows[i].cells[0], d, bold=True, bg="D9E2F3", align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in range(1, 5):
            fill(t.rows[i].cells[c], "○ / ✕", align=WD_ALIGN_PARAGRAPH.CENTER)

    para(doc, "", space_after=2)
    label_table(doc, [
        ("주 근무 가능일", "주        회        일   (    시간 이상)"),
        ("투입 가능일", "2026 년    월    일부터 가능"),
        ("희망 시급", "                    원 (경력 및 업종에 따라 협의)"),
        ("지원 직종", "□ 편의점  □ 카페·베이커리  □ 음식점·프랜차이즈  □ 홍보·리서치  □ 과외·학습  □ 물류·피킹  □ 기타 :"),
    ])

    para(doc, "", space_after=4)

    # career
    para(doc, "경력 (알바·아르바이트 중심)", size=11, bold=True, space_before=4, space_after=4, color=(0x44, 0x72, 0xC4))
    t = doc.add_table(rows=4, cols=5)
    style_table(t)
    set_col_widths(t, [3.0, 3.4, 4.6, 3.0, 2.2])
    header_row(t, ["기간", "근무지", "담당 업무", "퇴사 사유", "주 근무시간"])
    for r in range(1, 4):
        for c in range(5):
            fill(t.rows[r].cells[c], "")

    para(doc, "", space_after=4)

    # education
    para(doc, "학력", size=11, bold=True, space_before=4, space_after=4, color=(0x44, 0x72, 0xC4))
    t = doc.add_table(rows=3, cols=4)
    style_table(t)
    set_col_widths(t, [3.6, 5.4, 4.4, 2.8])
    header_row(t, ["기간", "학교명", "전공 / 상태", "비고"])
    for r in range(1, 3):
        for c in range(4):
            fill(t.rows[r].cells[c], "")

    para(doc, "", space_after=4)

    # skills
    para(doc, "보유 기술 / 자격증", size=11, bold=True, space_before=4, space_after=4, color=(0x44, 0x72, 0xC4))
    t = doc.add_table(rows=1, cols=1)
    style_table(t)
    fill(t.rows[0].cells[0],
         "· 예) POS기·카드단말기 사용 경험\n"
         "· 예) 요리기능사, 식품위생자격증, 운전면허(2종)\n"
         "· 예) 외국어 : 영어 회화 가능 / 일본어 기초\n"
         "· 예) PC : 한글, 엑셀, 포토샵 기초")

    para(doc, "", space_after=4)

    # self intro with alba example
    para(doc, "자기소개", size=11, bold=True, space_before=4, space_after=4, color=(0x44, 0x72, 0xC4))
    t = doc.add_table(rows=2, cols=1)
    style_table(t)
    fill(t.rows[0].cells[0],
         "예시) 안녕하세요. ○○점을 6개월간 운영하며 포스 관리, 재고 정리, 고객 응대를 맡았습니다. "
         "마감 후 재고를 매일 정리해 발주 실수를 줄였고, 아침 오픈 근무에 적응이 빠릅니다. "
         "주    회씩    시간 이상 꾸준히 근무 가능합니다.",
         size=9.5)
    fill(t.rows[1].cells[0], " (본인의 경험으로 바꿔서 4~6줄 작성)", color=(140, 140, 140), size=9.5)

    para(doc, "", space_after=6)
    para(doc, "작성일 : 2026 년    월    일        지원자 :                    (인) ",
         size=10, align=WD_ALIGN_PARAGRAPH.RIGHT)

    path = os.path.join(BASE, "resume", "알바이력서_단기알바_서식.docx")
    doc.save(path)
    return path


if __name__ == "__main__":
    print("OK:", make_alba_resume())
